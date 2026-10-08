import io
from datetime import date, datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app import excel, store
from backend.app.games import manual_id
from backend.app.main import app

client = TestClient(app)


def record(game_id, title, **over):
    mine = {
        "quantity": 1, "played": False, "rating": None, "review": None, "notes": None,
        "added_at": "2026-05-14T17:44:24", "purchase": {"date": None, "paid": None, "shop": None},
    }
    mine.update(over.pop("mine", {}))
    game = {
        "id": game_id, "title": title, "genres": ["보드게임"], "player_count": ["2-4인"], "play_time_minutes": 30,
        "price": 59000, "publisher": "코리아보드게임즈", "sale_link": None, "images": [], "tags": [],
        "source": "manual", "extra": {"sale_status": ["판매중"]}, "mine": mine,
    }
    game.update(over)
    return game


GAMES = [
    record("manual:a", "딕싯", mine={"played": True, "rating": 4, "purchase": {"date": "2026-10-01", "paid": 45000, "shop": "보드엠"}}),
    record("naver:1", "망령 열차", genres=["머더미스터리"], player_count=["5인"], price=None, play_time_minutes=None),
    record("manual:c", "카탄", mine={"quantity": 3}),
]


@pytest.fixture
def saved():
    store.save_collection(GAMES)
    return GAMES


def download() -> bytes:
    res = client.get("/api/excel")
    assert res.status_code == 200
    return res.content


def edit(data: bytes, cells: dict | None = None, new_rows: list[dict] | None = None, drop: tuple = ()) -> bytes:
    """cells: {(제목, 열 이름): 값}. new_rows: [{열 이름: 값}]. drop: 지울 행의 제목들."""
    wb = load_workbook(io.BytesIO(data))
    ws = wb[excel.SHEET]
    header = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    title_col = col["제목"]
    for (title, name), value in (cells or {}).items():
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, title_col).value == title)
        ws.cell(row, col[name], value)
    for title in drop:
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, title_col).value == title)
        ws.delete_rows(row)
    for values in new_rows or []:
        r = ws.max_row + 1
        for name, value in values.items():
            ws.cell(r, col[name], value)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def preview(data: bytes) -> dict:
    res = client.post("/api/excel/preview", content=data)
    assert res.status_code == 200, res.text
    return res.json()


def apply(data: bytes, base: str):
    return client.post("/api/excel/apply", params={"base": base}, content=data)


def by_id(game_id):
    return next(g for g in store.load_collection() if g["id"] == game_id)


# --- 내려받기 ---

def test_download_layout(saved):
    res = client.get("/api/excel")
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "filename*=UTF-8''" in res.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == [excel.SHEET, excel.GUIDE_SHEET]
    ws = wb[excel.SHEET]
    assert [c.value for c in ws[1]] == [c.header for c in excel.COLUMNS]
    rows = {r[1]: dict(zip([c.header for c in excel.COLUMNS], r)) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert list(rows) == ["딕싯", "망령 열차", "카탄"]  # 제목순
    dixit = rows["딕싯"]
    assert dixit["게임 ID"] == "manual:a" and dixit["인원"] == "2-4인" and dixit["해봤음"] == "O"
    assert dixit["구입일"] == datetime(2026, 10, 1) and dixit["구매가격(원)"] == 45000
    assert rows["망령 열차"]["정가(원)"] is None and rows["망령 열차"]["해봤음"] == "X"


def test_roundtrip_without_edits_changes_nothing(saved):
    result = preview(download())
    assert result["report"]["errors"] == []
    assert result["report"]["changes"] == []
    assert result["report"]["counts"] == {"new": 0, "updated": 0, "unchanged": 3, "not_in_file": 0}


# --- 고치기 ---

def test_edits_preview_then_apply(saved):
    data = edit(
        download(),
        {
            ("망령 열차", "정가(원)"): "35,000원",
            ("망령 열차", "인원"): "4-5인, 5인용",
            ("망령 열차", "해봤음"): "o",
            ("망령 열차", "장르"): "머더미스터리, 보드게임",
            ("카탄", "구입일"): "2025.12.24",
            ("카탄", "별점"): 5,
            ("카탄", "후기"): " 명작 ",
        },
    )
    result = preview(data)
    report = result["report"]
    assert report["errors"] == []
    assert report["counts"] == {"new": 0, "updated": 2, "unchanged": 1, "not_in_file": 0}
    murder = next(c for c in report["changes"] if c["id"] == "naver:1")
    assert {f["field"]: (f["before"], f["after"]) for f in murder["fields"]} == {
        "장르": ("머더미스터리", "보드게임, 머더미스터리"),
        "인원": ("5인", "4-5인"),
        "정가(원)": ("(빈칸)", "35,000"),
        "해봤음": ("X", "O"),
    }
    assert store.load_collection() == GAMES  # 미리보기는 저장하지 않는다

    res = apply(data, result["base"])
    assert res.status_code == 200, res.text
    game = by_id("naver:1")
    assert game["price"] == 35000 and game["player_count"] == ["4-5인"] and game["mine"]["played"] is True
    assert game["genres"] == ["보드게임", "머더미스터리"]
    catan = by_id("manual:c")
    assert catan["mine"]["purchase"]["date"] == "2025-12-24"
    assert catan["mine"]["rating"] == 5 and catan["mine"]["review"] == "명작"
    assert catan["extra"] == {"sale_status": ["판매중"]}  # 엑셀에 없는 칸은 그대로
    assert len(list(store.backup_dir().glob("collection-*.json"))) == 1


def test_blank_cell_keeps_value(saved):
    data = edit(download(), {("딕싯", "정가(원)"): None, ("딕싯", "후기"): "", ("딕싯", "구입일"): None})
    assert preview(data)["report"]["changes"] == []


def test_dash_clears_value(saved):
    data = edit(
        download(),
        {("딕싯", "정가(원)"): "-", ("딕싯", "별점"): "-", ("딕싯", "장르"): "-", ("딕싯", "구입일"): " - "},
    )
    result = preview(data)
    assert result["report"]["errors"] == []
    assert apply(data, result["base"]).status_code == 200
    game = by_id("manual:a")
    assert game["price"] is None and game["mine"]["rating"] is None and game["genres"] == []
    assert game["mine"]["purchase"]["date"] is None and game["mine"]["purchase"]["paid"] == 45000


@pytest.mark.parametrize("name", ["제목", "개수", "해봤음"])
def test_dash_not_allowed_on_required_fields(saved, name):
    errors = preview(edit(download(), {("딕싯", name): "-"}))["report"]["errors"]
    assert errors == [f"2행 {name}: 지울 수 없는 칸입니다"]


def test_new_row_without_id_adds_game(saved):
    data = edit(download(), new_rows=[{"제목": "스플렌더", "인원": "2-4인", "정가(원)": 65000, "개수": 2}])
    result = preview(data)
    assert result["report"]["counts"]["new"] == 1
    assert apply(data, result["base"]).status_code == 200
    game = by_id(manual_id("스플렌더"))
    assert game["price"] == 65000 and game["mine"]["quantity"] == 2 and game["mine"]["played"] is False


def test_removed_rows_are_kept(saved):
    data = edit(download(), drop=("카탄",))
    result = preview(data)
    assert result["report"]["counts"]["not_in_file"] == 1
    apply(data, result["base"])
    assert any(g["id"] == "manual:c" for g in store.load_collection())


def test_column_order_can_change(saved):
    wb = load_workbook(io.BytesIO(download()))
    ws = wb[excel.SHEET]
    ws.move_range("A1:A10", cols=20)  # 게임 ID 열을 맨 뒤로
    ws.delete_cols(1)
    out = io.BytesIO()
    wb.save(out)
    assert preview(out.getvalue())["report"]["errors"] == []


# --- 오류 ---

def test_row_errors_block_apply(saved):
    data = edit(
        download(),
        {
            ("딕싯", "별점"): 6,
            ("딕싯", "정가(원)"): "비쌈",
            ("망령 열차", "해봤음"): "?",
            ("카탄", "구입일"): "작년",
            ("카탄", "시간(분)"): 1.5,
        },
        new_rows=[
            {"인원": "4인"},  # 제목 없음
            {"제목": "딕 싯"},  # 띄어쓰기만 다른 같은 제목
            {"게임 ID": "manual:zzz", "제목": "없는 게임"},
            {"게임 ID": "manual:a", "제목": "딕싯"},  # ID 겹침
        ],
    )
    result = preview(data)
    assert result["report"]["errors"] == [
        "2행 정가(원): '비쌈'은(는) 숫자가 아닙니다",
        "3행 해봤음: O 또는 X로 적어 주세요",
        "4행 시간(분): 정수여야 합니다",
        "4행 구입일: 날짜(예: 2026-10-01)가 아닙니다",
        "5행: 제목을 입력해 주세요",
        "6행: 같은 제목의 게임이 이미 있습니다: 딕싯",
        "7행: 목록에 없는 게임 ID입니다 (manual:zzz). 새 게임이면 ID 칸을 비워 주세요",
        "8행: 게임 ID manual:a가 위의 행과 겹칩니다",
    ]
    res = apply(data, result["base"])
    assert res.status_code == 422
    assert store.load_collection() == GAMES


def test_range_errors_use_korean_messages(saved):
    errors = preview(edit(download(), {("딕싯", "별점"): 6, ("카탄", "개수"): 0}))["report"]["errors"]
    assert errors == ["2행 별점: 5 이하여야 합니다", "4행 개수: 1 이상이어야 합니다"]


def test_apply_refuses_when_list_changed_after_preview(saved):
    data = edit(download(), {("딕싯", "별점"): 5})
    base = preview(data)["base"]
    client.patch("/api/games/naver:1", json={"mine": {"notes": "화면에서 고침"}})
    res = apply(data, base)
    assert res.status_code == 409
    assert by_id("manual:a")["mine"]["rating"] == 4
    assert by_id("naver:1")["mine"]["notes"] == "화면에서 고침"


def test_apply_without_changes_does_not_save(saved):
    data = download()
    assert apply(data, preview(data)["base"]).status_code == 200
    assert not list(store.backup_dir().glob("collection-*.json"))


@pytest.mark.parametrize(
    "make, message",
    [
        (lambda: b"not an excel file", "읽을 수 없습니다"),
        (lambda: b"", "비어 있습니다"),
    ],
)
def test_unreadable_files(saved, make, message):
    res = client.post("/api/excel/preview", content=make())
    assert res.status_code == 422 and message in res.json()["detail"]


def test_missing_sheet_or_column(saved):
    wb = Workbook()
    wb.active.title = "다른 시트"
    out = io.BytesIO()
    wb.save(out)
    res = client.post("/api/excel/preview", content=out.getvalue())
    assert res.status_code == 422 and "시트가 없습니다" in res.json()["detail"]

    wb = load_workbook(io.BytesIO(download()))
    wb[excel.SHEET].delete_cols(3)  # 장르 열
    out = io.BytesIO()
    wb.save(out)
    res = client.post("/api/excel/preview", content=out.getvalue())
    assert res.status_code == 422 and "열이 없습니다: 장르" in res.json()["detail"]


def test_date_cells_and_text_numbers():
    assert excel._parse_cell(excel.BY_HEADER["구입일"], datetime(2026, 1, 2)) == "2026-01-02"
    assert excel._parse_cell(excel.BY_HEADER["구입일"], date(2026, 1, 2)) == "2026-01-02"
    for text in ["2026-01-02", "2026/1/2", "2026.1.2", "2026. 1. 2."]:
        assert excel._parse_cell(excel.BY_HEADER["구입일"], text) == "2026-01-02"
    for text in ["2026-13-01", "2026-02-30", "1월 2일"]:
        with pytest.raises(ValueError):
            excel._parse_cell(excel.BY_HEADER["구입일"], text)
    assert excel._parse_cell(excel.BY_HEADER["제목"], 1830.0) == "1830"
    assert excel._parse_cell(excel.BY_HEADER["정가(원)"], 59000.0) == 59000


def test_download_only_blanks(saved):
    res = client.get("/api/excel", params={"only": "blanks"})
    assert res.status_code == 200
    ws = load_workbook(io.BytesIO(res.content))[excel.SHEET]
    header = [c.value for c in ws[1]]
    rows = list(ws.iter_rows(min_row=2))
    assert [r[header.index("제목")].value for r in rows] == ["망령 열차"]  # 정가·시간이 빈 게임만
    fills = {header[i]: c.fill.start_color.rgb for i, c in enumerate(rows[0])}
    assert fills["정가(원)"].endswith("F8CBAD") and fills["시간(분)"].endswith("F8CBAD")
    assert not fills["인원"].endswith("F8CBAD")  # 인원은 있음
    assert "빈칸_" in res.headers["content-disposition"] or "%EB%B9%88%EC%B9%B8_" in res.headers["content-disposition"]


def test_blanks_file_upload_keeps_other_games(saved):
    data = edit(client.get("/api/excel", params={"only": "blanks"}).content, {("망령 열차", "시간(분)"): 120})
    result = preview(data)
    assert result["report"]["counts"] == {"new": 0, "updated": 1, "unchanged": 0, "not_in_file": 2}
    assert apply(data, result["base"]).status_code == 200
    assert by_id("naver:1")["play_time_minutes"] == 120
    assert len(store.load_collection()) == 3


def test_download_rejects_unknown_option(saved):
    assert client.get("/api/excel", params={"only": "x"}).status_code == 422


def test_old_paid_header_is_still_accepted(saved):
    # 열 이름을 '구매가격(원)'으로 바꾸기 전에 내려받은 엑셀도 올릴 수 있다
    wb = load_workbook(io.BytesIO(edit(download(), {("딕싯", "구매가격(원)"): 39000})))
    ws = wb[excel.SHEET]
    cell = next(c for c in ws[1] if c.value == "구매가격(원)")
    cell.value = "낸 가격(원)"
    buf = io.BytesIO()
    wb.save(buf)
    rows = excel.read_workbook(buf.getvalue())
    dixit = next(r for r in rows if r.cells.get("제목") == "딕싯")
    assert dixit.cells["구매가격(원)"] == 39000
