"""보유 목록 엑셀 내려받기·올리기. 파일을 읽고 쓰지 않는 순수 함수만 둔다 (저장은 main.py가 한다).

규칙 (엑셀의 '안내' 시트에도 같은 내용을 적는다)
- 행은 '게임 ID'로 맞춘다. ID가 빈 행에 제목이 있으면 새 게임으로 추가한다.
- 빈칸은 '그대로 두기'. 칸을 지우려면 '-' 하나만 적는다.
- 엑셀에서 지운 행의 게임은 목록에서 지우지 않는다 (삭제는 화면에서만).
- 한 행이라도 오류가 있으면 아무것도 적용하지 않는다 (고쳐서 다시 올린다).
"""
import io
import re
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from backend.app.games import GENRES, GameConflictError, GameIn, apply_patch, describe_error, new_game

SHEET = "보유 게임"
GUIDE_SHEET = "안내"
CLEAR = "-"
MAX_ROWS = 2000
TRUE_WORDS = {"o", "○", "y", "yes", "예", "네", "true", "1"}
FALSE_WORDS = {"x", "×", "n", "no", "아니오", "아니요", "false", "0"}


@dataclass(frozen=True)
class Col:
    header: str
    path: tuple[str, ...]  # 레코드 안의 위치. ("mine", "purchase", "date")
    kind: str  # id / text / list / int / bool / date
    width: int
    section: str  # id / game / mine
    clearable: bool = True


COLUMNS = [
    Col("게임 ID", ("id",), "id", 20, "id", False),
    Col("제목", ("title",), "text", 32, "game", False),
    Col("장르", ("genres",), "list", 16, "game"),
    Col("인원", ("player_count",), "list", 11, "game"),
    Col("시간(분)", ("play_time_minutes",), "int", 9, "game"),
    Col("정가(원)", ("price",), "int", 10, "game"),
    Col("제작사", ("publisher",), "text", 18, "game"),
    Col("판매 링크", ("sale_link",), "text", 28, "game"),
    Col("개수", ("mine", "quantity"), "int", 6, "mine", False),
    Col("해봤음", ("mine", "played"), "bool", 7, "mine", False),
    Col("별점", ("mine", "rating"), "int", 6, "mine"),
    Col("후기", ("mine", "review"), "text", 30, "mine"),
    Col("메모", ("mine", "notes"), "text", 24, "mine"),
    Col("구입일", ("mine", "purchase", "date"), "date", 12, "mine"),
    Col("낸 가격(원)", ("mine", "purchase", "paid"), "int", 11, "mine"),
    Col("산 곳", ("mine", "purchase", "shop"), "text", 14, "mine"),
]
BY_HEADER = {c.header: c for c in COLUMNS}


class ExcelReadError(Exception):
    """파일 자체를 읽을 수 없다 (엑셀이 아니거나 시트·열이 없음)."""


# ---------- 내려받기 ----------

FONT = "맑은 고딕"
FILLS = {
    "header": PatternFill("solid", start_color="3F3D56"),
    "id": PatternFill("solid", start_color="E7E6E6"),
    "game": PatternFill("solid", start_color="FFF2CC"),
    "mine": PatternFill("solid", start_color="E2EFDA"),
    "blank": PatternFill("solid", start_color="F8CBAD"),  # 채울 빈칸 (빈칸만 내려받기)
}
# 빈칸만 내려받기에서 채우라고 표시하는 칸
BLANK_TARGETS = ("인원", "시간(분)", "정가(원)")


def has_blanks(game: dict) -> bool:
    return not game["player_count"] or game["play_time_minutes"] is None or game["price"] is None
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _get(record: dict, path: tuple[str, ...]):
    for key in path:
        record = record.get(key) if isinstance(record, dict) else None
    return record


def _cell_value(col: Col, value):
    if col.kind == "list":
        return ", ".join(value or []) or None
    if col.kind == "bool":
        return "O" if value else "X"
    if col.kind == "date" and value:
        return date.fromisoformat(value)
    return value if value not in ("", []) else None


def build_workbook(games: list[dict], today: date | None = None, only_blanks: bool = False) -> bytes:
    """only_blanks면 인원·시간·정가 중 빈칸이 있는 게임만, 빈칸을 주황색으로 표시해서 내보낸다."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    for i, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=col.header)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = FILLS["header"]
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = col.width

    rows = sorted((g for g in games if not only_blanks or has_blanks(g)), key=lambda g: g["title"])
    for r, game in enumerate(rows, start=2):
        for i, col in enumerate(COLUMNS, start=1):
            value = _cell_value(col, _get(game, col.path))
            cell = ws.cell(row=r, column=i, value=value)
            cell.font = Font(name=FONT, color="7F7F7F" if col.section == "id" else "000000")
            cell.fill = FILLS["blank" if only_blanks and value is None and col.header in BLANK_TARGETS else col.section]
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=col.header in ("제목", "후기", "메모"))
            if col.kind == "int" and col.header.endswith("(원)"):
                cell.number_format = "#,##0"
            if col.kind == "date":
                cell.number_format = "yyyy-mm-dd"

    last = len(rows) + 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    _add_checks(ws, last + 300)
    notes = {
        "게임 ID": "행을 찾는 데 씁니다. 바꾸지 마세요. 새 게임은 이 칸을 비우고 제목을 적으면 추가됩니다.",
        "장르": f"{', '.join(GENRES)} 중에서 (쉼표로 구분, 둘 다 가능)",
        "인원": "예: 4인 / 2-4인 / 5+gm (게임마스터 필요). 여러 개면 쉼표로.",
        "해봤음": "O = 해봤음, X = 안 해봄",
    }
    for i, col in enumerate(COLUMNS, start=1):
        if col.header in notes:
            ws.cell(row=1, column=i).comment = Comment(notes[col.header], "boardgame-home")

    _add_guide(wb, len(rows), today or date.today(), only_blanks)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _add_checks(ws, until: int) -> None:
    """엑셀에서 바로 잡아 주는 입력 검사 (빈칸과 '-'는 허용). 서버도 올릴 때 한 번 더 검사한다."""
    letter = {c.header: get_column_letter(i) for i, c in enumerate(COLUMNS, start=1)}

    def add(header: str, formula: str, message: str) -> None:
        col = letter[header]
        dv = DataValidation(type="custom", formula1=formula.replace("@", f"{col}2"), allow_blank=True)
        dv.error, dv.errorTitle, dv.showErrorMessage = message, "입력 형식", True
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{until}")

    number = 'OR(@="-",AND(ISNUMBER(@),@=INT(@),@>={low}))'
    add("시간(분)", number.format(low=1), "시간은 1 이상의 숫자(분)로 적거나, 지우려면 - 를 적어 주세요.")
    add("정가(원)", number.format(low=0), "정가는 0 이상의 숫자(원)로 적거나, 지우려면 - 를 적어 주세요.")
    add("낸 가격(원)", number.format(low=0), "낸 가격은 0 이상의 숫자(원)로 적거나, 지우려면 - 를 적어 주세요.")
    add("개수", "AND(ISNUMBER(@),@=INT(@),@>=1)", "개수는 1 이상의 숫자로 적어 주세요.")
    add("별점", 'OR(@="-",AND(ISNUMBER(@),@=INT(@),@>=1,@<=5))', "별점은 1~5로 적거나, 지우려면 - 를 적어 주세요.")
    add("해봤음", 'OR(@="O",@="X")', "해봤음은 O 또는 X로 적어 주세요.")


def _add_guide(wb: Workbook, count: int, today: date, only_blanks: bool = False) -> None:
    ws = wb.create_sheet(GUIDE_SHEET)
    blanks = [
        ("빈칸 채우기용 파일", True),
        ("인원·시간·정가 중 빈칸이 있는 게임만 담았습니다. 채울 빈칸은 주황색입니다.", False),
        ("게임 상자나 공식 판매처에서 확인한 값만 적고, 모르면 빈칸으로 두세요. 정가는 할인 전 가격입니다.", False),
        ("이 파일에 없는 게임은 올려도 그대로입니다.", False),
        ("", False),
    ] if only_blanks else []
    lines = [
        ("보유 게임 엑셀 편집 안내", True),
        (f"내려받은 날: {today.isoformat()}  ·  게임 {count}개", False),
        ("", False),
        *blanks,
        ("쓰는 법", True),
        ("이 파일을 고친 뒤 사이트의 [엑셀 올리기]로 올리면, 바뀔 내용을 먼저 보여 주고 [적용]을 눌러야 저장됩니다.", False),
        ("저장하기 전 상태는 자동으로 백업됩니다.", False),
        ("", False),
        ("칸 색깔", True),
        ("회색 — 게임 ID. 행을 찾는 데 쓰니 바꾸지 마세요.", False),
        ("노란색 — 게임 정보 (제목·장르·인원·시간·정가·제작사·판매 링크)", False),
        ("초록색 — 내 정보 (개수·해봤음·별점·후기·메모·구입 정보)", False),
        ("", False),
        ("빈칸과 지우기", True),
        ("빈칸은 '그대로 두기'입니다. 모르는 값은 빈칸으로 두세요.", False),
        ("칸의 값을 지우려면 - (빼기 하나)만 적으세요. 제목·개수·해봤음은 지울 수 없습니다.", False),
        ("", False),
        ("적는 형식", True),
        (f"장르: {', '.join(GENRES)} (쉼표로 구분, 둘 다 가능)", False),
        ("인원: 4인 / 2-4인 / 5+gm (게임마스터 필요). 여러 개면 쉼표로", False),
        ("시간: 분 단위 숫자 · 정가·낸 가격: 원 단위 숫자 (정가는 할인 전 가격) · 별점: 1~5", False),
        ("해봤음: O 또는 X · 구입일: 2026-10-01 형식", False),
        ("", False),
        ("행 추가·삭제", True),
        ("새 게임: 맨 아래 행에 게임 ID는 비우고 제목부터 적으세요. 같은 제목(띄어쓰기 무시)이 이미 있으면 오류입니다.", False),
        ("행을 지워도 사이트에서는 지워지지 않습니다. 게임 삭제는 사이트의 상세 창에서 하세요.", False),
        ("열 이름과 시트 이름은 바꾸지 마세요. 열 순서는 바꿔도 됩니다.", False),
    ]
    for r, (text, bold) in enumerate(lines, start=1):
        ws.cell(row=r, column=1, value=text).font = Font(name=FONT, bold=bold, size=13 if r == 1 else 11)
    ws.column_dimensions["A"].width = 110


# ---------- 올리기: 읽기 ----------

@dataclass
class Row:
    number: int  # 엑셀 행 번호
    game_id: str | None
    cells: dict[str, object]  # 열 이름 → 칸 값 (빈칸은 넣지 않음)


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def read_workbook(data: bytes) -> list[Row]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # 손상된 파일, 다른 형식
        raise ExcelReadError("엑셀(.xlsx) 파일을 읽을 수 없습니다.") from e
    if SHEET not in wb.sheetnames:
        raise ExcelReadError(f"'{SHEET}' 시트가 없습니다. 사이트에서 내려받은 파일을 고쳐서 올려 주세요.")
    values = wb[SHEET].iter_rows(values_only=True)
    header = [str(v).strip() if v is not None else "" for v in next(values, ())]
    missing = [c.header for c in COLUMNS if c.header not in header]
    if missing:
        raise ExcelReadError(f"열이 없습니다: {', '.join(missing)}. 열 이름을 바꾸지 말아 주세요.")
    index = {h: header.index(h) for h in BY_HEADER}

    rows = []
    for number, raw in enumerate(values, start=2):
        if number > MAX_ROWS + 1:
            raise ExcelReadError(f"행이 너무 많습니다 ({MAX_ROWS}행까지).")
        raw = list(raw) + [None] * (len(header) - len(raw))
        cells = {h: raw[i] for h, i in index.items() if not _blank(raw[i])}
        if not cells:
            continue  # 빈 행
        game_id = cells.pop("게임 ID", None)
        rows.append(Row(number, str(game_id).strip() if game_id is not None else None, cells))
    return rows


# ---------- 올리기: 칸 값 → 수정 값 ----------

def _parse_cell(col: Col, value):
    """칸 값 하나 → 저장할 값. 형식이 틀리면 ValueError (범위 검사는 GameIn이 한다)."""
    if isinstance(value, str):
        value = value.strip()
    if value == CLEAR:
        if not col.clearable:
            raise ValueError("지울 수 없는 칸입니다")
        return [] if col.kind == "list" else None
    if col.kind == "text":
        # 엑셀이 숫자로 읽은 제목(예: 1830)이 "1830.0"이 되지 않게
        return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
    if col.kind == "list":
        return [part.strip() for part in re.split(r"[,、/]", str(value)) if part.strip()]
    if col.kind == "int":
        if isinstance(value, bool):
            raise ValueError("숫자가 아닙니다")
        if isinstance(value, (int, float)):
            number = value
        else:
            try:
                number = float(str(value).replace(",", "").replace("원", "").replace("분", "").strip())
            except ValueError:
                raise ValueError(f"'{value}'은(는) 숫자가 아닙니다") from None
        if number != int(number):
            raise ValueError("정수여야 합니다")
        return int(number)
    if col.kind == "bool":
        word = str(value).strip().lower()
        if word in TRUE_WORDS:
            return True
        if word in FALSE_WORDS:
            return False
        raise ValueError("O 또는 X로 적어 주세요")
    if col.kind == "date":
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        # 2026-10-01, 2026.10.1, 2026/1/2 모두 받는다
        match = re.fullmatch(r"(\d{4})\s*[-./]\s*(\d{1,2})\s*[-./]\s*(\d{1,2})\.?", str(value).strip())
        try:
            if not match:
                raise ValueError
            return date(*map(int, match.groups())).isoformat()
        except ValueError:  # 형식이 다르거나 없는 날짜(2월 30일 등)
            raise ValueError("날짜(예: 2026-10-01)가 아닙니다") from None
    raise AssertionError(col.kind)


def _set(target: dict, path: tuple[str, ...], value) -> None:
    for key in path[:-1]:
        target = target.setdefault(key, {})
    target[path[-1]] = value


def _row_input(row: Row) -> tuple[GameIn | None, list[str]]:
    patch: dict = {}
    errors = []
    for header, value in row.cells.items():
        col = BY_HEADER[header]
        try:
            _set(patch, col.path, _parse_cell(col, value))
        except ValueError as e:
            errors.append(f"{header}: {e}")
    if errors:
        return None, errors
    try:
        return GameIn(**patch), []
    except ValidationError as e:
        return None, [describe_error(err) for err in e.errors()]


# ---------- 올리기: 바뀔 내용 ----------

def _show(col: Col, value) -> str:
    if value is None or value == []:
        return "(빈칸)"
    if col.kind == "list":
        return ", ".join(value)
    if col.kind == "bool":
        return "O" if value else "X"
    if col.kind == "int" and col.header.endswith("(원)"):
        return f"{value:,}"
    return str(value)


def _diff(before: dict, after: dict) -> list[dict]:
    return [
        {"field": c.header, "before": _show(c, _get(before, c.path)), "after": _show(c, _get(after, c.path))}
        for c in COLUMNS[1:]
        if _get(before, c.path) != _get(after, c.path)
    ]


def plan_import(games: list[dict], rows: list[Row]) -> tuple[list[dict], dict]:
    """(반영한 새 목록, 보고서). 보고서의 errors가 비어 있지 않으면 적용하면 안 된다."""
    current = [dict(g) for g in games]
    original = {g["id"]: g for g in games}
    changes, errors, seen = [], [], set()
    unchanged = 0

    for row in rows:
        where = f"{row.number}행"
        if row.game_id is not None:
            if row.game_id in seen:
                errors.append(f"{where}: 게임 ID {row.game_id}가 위의 행과 겹칩니다")
                continue
            seen.add(row.game_id)
            if row.game_id not in original:
                errors.append(f"{where}: 목록에 없는 게임 ID입니다 ({row.game_id}). 새 게임이면 ID 칸을 비워 주세요")
                continue
        data, problems = _row_input(row)
        if problems:
            errors.extend(f"{where} {p}" for p in problems)
            continue
        try:
            if row.game_id is None:
                game = new_game(current, data)
                current.append(game)
                changes.append({"row": row.number, "id": game["id"], "title": game["title"], "kind": "new",
                                "fields": _diff({}, game)})
            else:
                before = next(g for g in current if g["id"] == row.game_id)
                current = apply_patch(current, row.game_id, data)
                after = next(g for g in current if g["id"] == row.game_id)
                fields = _diff(before, after)
                if fields:
                    changes.append({"row": row.number, "id": row.game_id, "title": after["title"], "kind": "update",
                                    "fields": fields})
                else:
                    unchanged += 1
        except GameConflictError as e:
            errors.append(f"{where}: {e}")
        except ValueError as e:  # 제목 없는 새 행
            errors.append(f"{where}: {e}")

    report = {
        "changes": changes,
        "errors": errors,
        "counts": {
            "new": sum(c["kind"] == "new" for c in changes),
            "updated": sum(c["kind"] == "update" for c in changes),
            "unchanged": unchanged,
            "not_in_file": len(set(original) - seen),
        },
    }
    return current, report
